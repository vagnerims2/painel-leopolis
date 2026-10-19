"""Gera a planilha Excel do painel a partir dos CSVs do build.py.
Uso: python3 gerar_planilha.py <pasta_saida_do_build> <arquivo.xlsx>
"""
import csv, os, sys
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SAIDA, XLSX = sys.argv[1], sys.argv[2]
ler = lambda n: list(csv.DictReader(open(os.path.join(SAIDA, n), encoding='utf-8')))
f = lambda x: float(x or 0)
MOEDA, QTD, INT, DATA = '#,##0.00', '#,##0.###', '0', 'DD/MM/YYYY'
AZUL = PatternFill('solid', fgColor='1B2A47')

def para_data(v):
    try:
        return datetime.strptime(v, '%Y-%m-%d')
    except (ValueError, TypeError):
        return v or None

def aba(wb, nome, linhas, colunas, tipos, larguras=None):
    """colunas: [(campo, titulo)]; tipos: {campo: 'moeda'|'qtd'|'int'|'data'}"""
    ws = wb.create_sheet(nome)
    ws.append([t for _, t in colunas])
    for r in linhas:
        out = []
        for c, _ in colunas:
            v = r.get(c, '')
            t = tipos.get(c)
            if t in ('moeda', 'qtd'):
                v = f(v)
            elif t == 'int':
                v = int(f(v)) if str(v).strip() not in ('', 'None') else None
            elif t == 'data':
                v = para_data(v)
            out.append(v)
        ws.append(out)
    for j, (c, _) in enumerate(colunas, 1):
        fmt = {'moeda': MOEDA, 'qtd': QTD, 'int': INT, 'data': DATA}.get(tipos.get(c))
        if fmt:
            for cel in ws.iter_cols(min_col=j, max_col=j, min_row=2, max_row=ws.max_row):
                for x in cel:
                    x.number_format = fmt
        ws.column_dimensions[get_column_letter(j)].width = (larguras or {}).get(c, 14)
    if ws.max_row > 1:
        ref = f"A1:{get_column_letter(len(colunas))}{ws.max_row}"
        t = Table(displayName='T_' + ''.join(ch for ch in nome if ch.isalnum()), ref=ref)
        t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(t)
    ws.freeze_panes = 'A2'
    return ws

E, IT, L, D, F, M = (ler(n) for n in ('empenhos.csv', 'itens.csv', 'liquidacoes.csv', 'dotacoes.csv', 'saldo_fontes.csv', 'movimento.csv'))
wb = Workbook(); res = wb.active; res.title = 'Resumo'

# ---- resumo
ult = max((r['data'] for r in M), default='')
par = [r for r in E if r['tipo'] == 'Ordinário' and f(r['liquidado']) == 0 and f(r['saldo_a_pagar']) > 0 and int(f(r['dias'])) >= 10]
dm = [r for r in D if r['recorrente'] == 'Sim' and f(r['falta_fechar_mes']) > 0]
res['A1'] = 'Painel de Empenhos — Prefeitura Municipal de Leópolis'; res['A1'].font = Font(bold=True, size=14, color='FFFFFF'); res['A1'].fill = AZUL
res.merge_cells('A1:C1')
res['A2'] = f"Dados do Portal da Transparência até {'/'.join(reversed(ult.split('-')))} · planilha gerada em {datetime.now().strftime('%d/%m/%Y %H:%M')}"
linhas = [('Empenhado no ano (líquido)', sum(f(r['empenhado_liquido']) for r in M), MOEDA),
          ('Liquidado no ano (inclui restos a pagar)', sum(f(r['liquidado']) for r in M), MOEDA),
          ('Pago no ano (inclui restos a pagar)', sum(f(r['pago']) for r in M), MOEDA),
          ('Saldo a pagar (oficial, por fonte)', sum(f(r['saldo_a_pagar']) for r in F), MOEDA),
          ('   falta liquidar', sum(f(r['falta_liquidar']) for r in E), MOEDA),
          ('   liquidado e não pago', sum(f(r['liquidado_nao_pago']) for r in E), MOEDA),
          ('Empenhos no ano', len(E), INT),
          ('Empenhos com saldo a pagar', sum(1 for r in E if f(r['saldo_a_pagar']) > 0), INT),
          ('Empenhos ordinários sem liquidação há 10+ dias', len(par), INT),
          ('   dos quais há 30+ dias', sum(1 for r in par if int(f(r['dias'])) >= 30), INT),
          ('Dotações recorrentes que não fecham o mês', len(dm), INT),
          ('   falta para fechar o mês', sum(f(r['falta_fechar_mes']) for r in dm), MOEDA)]
for i, (t, v, fmt) in enumerate(linhas, 4):
    res.cell(i, 1, t); c = res.cell(i, 2, v); c.number_format = fmt
    if not t.startswith('   '):
        res.cell(i, 1).font = Font(bold=True)
res.column_dimensions['A'].width = 48; res.column_dimensions['B'].width = 20
res.cell(len(linhas) + 6, 1, 'As outras abas têm filtro em cada coluna (setinha no cabeçalho). Fonte: Portal da Transparência de Leópolis (Equiplano); em caso de divergência, vale o sistema oficial.').alignment = Alignment(wrap_text=False)

aba(wb, 'Empenhos', E,
    [('empenho', 'Empenho'), ('data', 'Data'), ('tipo', 'Tipo'), ('fornecedor', 'Fornecedor'), ('secretaria', 'Secretaria'), ('modalidade', 'Modalidade'), ('licitacao', 'Licitação'),
     ('contrato', 'Contrato'), ('natureza', 'Natureza'), ('subelemento', 'Código da despesa'), ('tipo_despesa', 'Tipo de despesa'), ('dotacao', 'Dotação'), ('fonte', 'Fonte'),
     ('valor', 'Valor empenhado'), ('anulado', 'Anulado'), ('liquidado', 'Liquidado'), ('pago', 'Pago'), ('falta_liquidar', 'Falta liquidar'), ('liquidado_nao_pago', 'Liquidado não pago'),
     ('saldo_a_pagar', 'Saldo a pagar'), ('situacao', 'Situação'), ('estornos', 'Estornos / anulações'), ('dias', 'Dias'), ('primeira_liquidacao', '1ª liquidação'), ('dias_ate_liquidar', 'Dias até liquidar'),
     ('ultimo_pagamento', 'Último pagamento'), ('dias_ate_pagar', 'Dias até pagar'), ('produtos', 'Produtos / histórico')],
    {'empenho': 'int', 'data': 'data', 'valor': 'moeda', 'anulado': 'moeda', 'liquidado': 'moeda', 'pago': 'moeda', 'falta_liquidar': 'moeda', 'liquidado_nao_pago': 'moeda',
     'saldo_a_pagar': 'moeda', 'dias': 'int', 'primeira_liquidacao': 'data', 'dias_ate_liquidar': 'int', 'ultimo_pagamento': 'data', 'dias_ate_pagar': 'int'},
    {'fornecedor': 38, 'secretaria': 30, 'tipo_despesa': 34, 'fonte': 30, 'produtos': 60, 'estornos': 50, 'subelemento': 17, 'licitacao': 18})
aba(wb, 'Itens', IT,
    [('empenho', 'Empenho'), ('data', 'Data'), ('fornecedor', 'Fornecedor'), ('secretaria', 'Secretaria'), ('licitacao', 'Licitação'), ('produto', 'Produto'), ('marca', 'Marca'),
     ('quantidade', 'Quantidade'), ('valor_unitario', 'Valor unitário'), ('valor_total', 'Valor total'), ('situacao', 'Situação do empenho')],
    {'empenho': 'int', 'data': 'data', 'quantidade': 'qtd', 'valor_unitario': 'moeda', 'valor_total': 'moeda'},
    {'fornecedor': 36, 'secretaria': 30, 'produto': 50, 'marca': 18, 'licitacao': 18})
aba(wb, 'Liquidações', L,
    [('liquidacao', 'Liquidação'), ('data', 'Data'), ('empenho', 'Empenho'), ('fornecedor', 'Fornecedor'), ('secretaria', 'Secretaria'), ('tipo_despesa', 'Tipo de despesa'),
     ('fonte', 'Fonte'), ('valor', 'Valor'), ('documento', 'Documento fiscal'), ('vencimento', 'Vencimento'), ('pago', 'Paga')],
    {'liquidacao': 'int', 'data': 'data', 'valor': 'moeda', 'vencimento': 'data'},
    {'fornecedor': 36, 'secretaria': 30, 'tipo_despesa': 32, 'fonte': 30, 'documento': 22, 'empenho': 16})
aba(wb, 'Dotações', D,
    [('dotacao', 'Dotação'), ('secretaria', 'Secretaria'), ('subelemento', 'Código da despesa'), ('tipo_despesa', 'Tipo de despesa'), ('fonte', 'Fonte'), ('saldo_atual', 'Saldo atual'),
     ('ultimo_empenho', 'Último empenho'), ('data_ultimo', 'Data do último'), ('empenhado_ano', 'Empenhado no ano'), ('media_mensal', 'Média mensal (mediana 6 meses)'),
     ('empenhado_mes', 'Empenhado no mês'), ('recorrente', 'Recorrente'), ('falta_fechar_mes', 'Falta p/ fechar o mês'), ('precisa_proximo_mes', 'Precisa p/ o mês que vem'),
     ('total_suplementar', 'Total a suplementar'), ('meses', 'Últimos 6 meses')],
    {'saldo_atual': 'moeda', 'ultimo_empenho': 'int', 'data_ultimo': 'data', 'empenhado_ano': 'moeda', 'media_mensal': 'moeda', 'empenhado_mes': 'moeda',
     'falta_fechar_mes': 'moeda', 'precisa_proximo_mes': 'moeda', 'total_suplementar': 'moeda'},
    {'secretaria': 30, 'tipo_despesa': 34, 'fonte': 30, 'meses': 60, 'subelemento': 17})
aba(wb, 'Saldo por fonte', F, [('fonte', 'Fonte'), ('codigo', 'Código'), ('saldo_a_pagar', 'Saldo a pagar')], {'saldo_a_pagar': 'moeda'}, {'fonte': 60})
aba(wb, 'Movimento diário', sorted(M, key=lambda r: r['data']),
    [('data', 'Data'), ('empenhado', 'Empenhado'), ('estornado', 'Estornado'), ('empenhado_liquido', 'Empenhado líquido'), ('liquidado', 'Liquidado'), ('pago', 'Pago')],
    {'data': 'data', 'empenhado': 'moeda', 'estornado': 'moeda', 'empenhado_liquido': 'moeda', 'liquidado': 'moeda', 'pago': 'moeda'})
os.makedirs(os.path.dirname(os.path.abspath(XLSX)), exist_ok=True)
wb.save(XLSX)
print('ok', XLSX)
